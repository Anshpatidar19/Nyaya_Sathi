"""One extraction pipeline for every uploaded document.

    upload -> detect type (from the bytes, not the client's label)
           -> text where the file has text  (txt, DOCX, text-layer PDF pages)
           -> Gemini OCR where it doesn't   (photos, scans, handwriting,
                                             scanned pages inside a PDF)
           -> ONE normalised text, page-marked
           -> the existing analysis pipelines, unchanged

Everything downstream - Ask with a document, Review, Arguments, Document
Intelligence in chat threads - calls load() and gets plain text back, the
same as before. None of them know or care whether the text came from a
Word file or a photographed handwritten application.

Latency
-------
OCR is a Gemini call, seconds per page. It must not land on the question
path, so:

  1. prewarm() is started the moment a file is uploaded. The upload
     response does not wait for it. By the time the user has typed their
     question the text is usually already extracted.
  2. load() checks, in order: memory cache -> the prewarm still running
     (awaits only what is left of it) -> a saved extraction in storage ->
     extract from scratch. A text PDF or DOCX is also prewarmed, which moves
     pdfplumber off the first question too.
  3. The result is saved next to the file in the PRIVATE bucket as
     <path>.extract.json, so a restart or a second worker never pays for
     OCR again. Same user folder, same access rules as the file itself.
  4. Only the pages that need OCR are sent. A 40-page typed PDF with two
     scanned annexure pages OCRs two pages. Batches run concurrently.

Honesty
-------
The OCR prompt forbids guessing: unreadable words come back as [illegible].
If too much of a document is illegible, load() refuses with a plain
message asking for a clearer scan, rather than handing the analysis step a
transcript that is mostly invention. Partly legible documents go through
with a note at the top of the text, which the analysis prompts see.
"""

from __future__ import annotations

import asyncio
import io
import json
import logging
import re
import shutil
import subprocess
import tempfile
from collections import OrderedDict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import httpx

from . import gemini, storage
from .config import settings

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Limits
# ---------------------------------------------------------------------------
# Upload size stays storage.MAX_BYTES (10 MB). 10 MB base64-encodes to
# ~13.4 MB, inside Gemini's 20 MB inline request limit.

MAX_OCR_PAGES = 30      # pages sent to OCR per document; the rest are noted
OCR_BATCH_PAGES = 4     # pages per Gemini call
OCR_CONCURRENCY = 4     # batches in flight at once
OCR_TOKENS_PER_PAGE = 1800

# A PDF page with fewer real characters than this in its text layer is
# treated as scanned. Scans often carry a stray header or page number in a
# text layer, so "any text at all" is not the right test.
MIN_TEXT_CHARS = 40

# Refuse rather than analyse when more than this share of the transcribed
# words are [illegible].
MAX_ILLEGIBLE_SHARE = 0.35

EXTRACT_VERSION = 1     # bump to invalidate saved extractions


# ---------------------------------------------------------------------------
# Result
# ---------------------------------------------------------------------------

@dataclass
class Extraction:
    status: str                      # "ok" | "unreadable"
    text: str = ""
    method: str = ""                 # txt | docx | doc | pdf-text | pdf-ocr | pdf-mixed | image-ocr
    pages: int = 1
    ocr_pages: List[int] = field(default_factory=list)
    unclear_pages: List[int] = field(default_factory=list)
    handwritten: bool = False
    message: str = ""                # why it is unreadable, for the user
    version: int = EXTRACT_VERSION


class Unreadable(ValueError):
    """The file was read, and it cannot be read well enough to analyse.

    A ValueError so every existing caller's `except ValueError -> 400`
    already handles it with this message.
    """


# ---------------------------------------------------------------------------
# Type detection
# ---------------------------------------------------------------------------
# The browser's Content-Type is a claim. The first bytes are evidence. A
# renamed executable claiming to be a PDF is rejected at upload.

IMAGE_MIMES = {"image/jpeg", "image/png", "image/webp", "image/heic", "image/heif"}


def detect_kind(data: bytes, content_type: str = "", filename: str = "") -> Optional[str]:
    """pdf | docx | doc | image/<x> | txt, or None if it is none of them."""
    head = data[:16]
    name = (filename or "").lower()
    ct = (content_type or "").lower()

    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if head[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[4:8] == b"ftyp" and data[8:12] in (b"heic", b"heix", b"heif", b"mif1", b"msf1", b"hevc"):
        return "image/heic"
    if head.startswith(b"PK\x03\x04"):
        # DOCX is a zip; so is every other Office Open XML file and any zip.
        if b"word/" in data[:4000] or name.endswith(".docx") or "wordprocessingml" in ct:
            return "docx"
        return None
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):   # OLE: legacy .doc
        return "doc" if (name.endswith(".doc") or "msword" in ct) else None
    if ct == "text/plain" or name.endswith(".txt"):
        try:
            data[:4000].decode("utf-8")
            return "txt"
        except UnicodeDecodeError:
            return None
    return None


# ---------------------------------------------------------------------------
# Local extractors
# ---------------------------------------------------------------------------

def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _join_pages(pages: List[str]) -> str:
    """Page markers only when there is more than one page to tell apart.

    They let an answer say "page 3" and let a reader check it.
    """
    if len(pages) == 1:
        return _clean(pages[0])
    return "\n\n".join(f"[Page {i}]\n{_clean(t)}" for i, t in enumerate(pages, 1) if _clean(t))


def _docx_text(data: bytes) -> str:
    """Paragraphs AND tables, in document order.

    python-docx's document.paragraphs skips tables entirely, and schedules
    of payment, rent tables and party details in agreements live in tables.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    out: List[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append(Paragraph(child, document).text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells, seen = [], set()
                for c in row.cells:           # merged cells repeat; keep once
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc))
                        cells.append(c.text.strip())
                if any(cells):
                    out.append(" | ".join(cells))
    return _clean("\n".join(out))


def _doc_text(data: bytes) -> str:
    """Legacy binary .doc (Word 97-2003).

    Converted with LibreOffice when it is installed, which is exact. Without
    it, a best-effort read of the text runs stored in the file, which works
    for ordinary letters and agreements. If neither gives real text, the
    user is asked for DOCX or PDF - there is no Gemini input type for .doc.
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            with tempfile.TemporaryDirectory() as tmp:
                src = Path(tmp) / "in.doc"
                src.write_bytes(data)
                subprocess.run(
                    [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)],
                    check=True, timeout=45, capture_output=True,
                )
                out = Path(tmp) / "in.docx"
                if out.exists():
                    return _docx_text(out.read_bytes())
        except Exception as exc:
            logger.warning("LibreOffice .doc conversion failed (%s); using fallback", exc)

    candidates = []
    for enc, pat in (
        ("utf-16-le", r"[\x20-\x7E\u0900-\u097F\u2018-\u201D\n\r\t]{20,}"),
        ("cp1252", r"[\x20-\x7E\n\r\t]{20,}"),
    ):
        decoded = data.decode(enc, errors="ignore")
        runs = [r for r in re.findall(pat, decoded) if " " in r]
        candidates.append("\n".join(runs))
    best = max(candidates, key=lambda t: sum(ch.isalpha() for ch in t))
    if sum(ch.isalpha() for ch in best) < 200:
        raise ValueError(
            "This older Word (.doc) file could not be read. Save it as .docx "
            "or PDF and upload it again."
        )
    return _clean(best)


def _pdf_pages(data: bytes) -> List[str]:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return [p.extract_text() or "" for p in pdf.pages]


def _pdf_subset(data: bytes, page_numbers: List[int]) -> bytes:
    """A new PDF holding only the given 1-based pages.

    pypdfium2 ships with pdfplumber, so this adds no dependency. Sending
    only the scanned pages is what keeps a mostly-typed PDF cheap and fast.
    """
    import pypdfium2 as pdfium
    src = pdfium.PdfDocument(data)
    try:
        dst = pdfium.PdfDocument.new()
        dst.import_pages(src, [n - 1 for n in page_numbers])
        buf = io.BytesIO()
        dst.save(buf)
        return buf.getvalue()
    finally:
        src.close()


def _real_chars(text: str) -> int:
    return len(re.sub(r"\s+", "", text or ""))


# ---------------------------------------------------------------------------
# Gemini OCR
# ---------------------------------------------------------------------------

_OCR_SYSTEM = """You transcribe legal documents from scans, photos and \
handwriting for an Indian legal-information platform. You are an OCR \
engine, not an assistant.

Rules:
- Transcribe EXACTLY what is written. Do not summarise, translate, correct \
spelling, or complete sentences. Keep the original language and script \
(Hindi stays in Devanagari, English in English, mixed stays mixed).
- NEVER GUESS. A word or number you cannot read with confidence is written \
as [illegible]. A whole unreadable line is [illegible line]. Guessing a \
name, date, amount or section number is worse than leaving a gap, because \
people act on these documents.
- Preserve structure: headings on their own line, numbered and lettered \
clauses kept as numbered, line breaks between paragraphs, tables as rows \
with cells separated by " | ".
- Non-text marks: [signature], [stamp: <text on the stamp if readable>], \
[seal], [photo], [thumb impression], [strikethrough: <text>] for crossed-out \
text that is still readable.
- Printed form fields filled by hand: write "<printed label>: <handwritten \
value>".
- "legibility" per page: "clear" (all or nearly all readable), "partial" \
(readable overall, some [illegible]), "illegible" (too unclear to transcribe \
meaningfully, or not a document at all).

Return ONLY JSON:
{
  "pages": [
    {"page": 1, "legibility": "clear|partial|illegible", "text": "..."}
  ],
  "handwritten": true or false
}"""


def _ocr_model() -> Optional[str]:
    return settings.gemini_ocr_model or None


async def _ocr_call(mime: str, blob: bytes, n_pages: int) -> Dict[str, Any]:
    what = "this image" if mime.startswith("image/") else f"all {n_pages} page(s) of this PDF, in order"
    prompt = (
        f"Transcribe {what}. Return one entry in \"pages\" per page"
        f"{'' if mime.startswith('image/') else ', numbered 1 to ' + str(n_pages)}."
    )
    return await gemini.generate_json(
        prompt,
        _OCR_SYSTEM,
        temperature=0.0,
        max_output_tokens=min(OCR_TOKENS_PER_PAGE * max(n_pages, 1) + 400, 16000),
        thinking_budget=0,            # reading, not reasoning; no thinking wait
        media=[(mime, blob)],
        model=_ocr_model(),
    )


def _pages_from(data: Dict[str, Any], expected: int) -> Tuple[List[Tuple[str, str]], bool]:
    """(text, legibility) per page, padded/trimmed to what was sent."""
    raw = data.get("pages") or []
    out: List[Tuple[str, str]] = []
    for i in range(expected):
        item = raw[i] if i < len(raw) and isinstance(raw[i], dict) else {}
        leg = str(item.get("legibility") or "partial").lower()
        if leg not in ("clear", "partial", "illegible"):
            leg = "partial"
        out.append((str(item.get("text") or ""), leg))
    return out, bool(data.get("handwritten"))


async def _ocr_image(mime: str, data: bytes) -> Tuple[List[Tuple[str, str]], bool]:
    return _pages_from(await _ocr_call(mime, data, 1), 1)


async def _ocr_pdf_pages(data: bytes, page_numbers: List[int]) -> Tuple[Dict[int, Tuple[str, str]], bool]:
    """OCR the given 1-based pages, in concurrent batches."""
    batches = [page_numbers[i:i + OCR_BATCH_PAGES] for i in range(0, len(page_numbers), OCR_BATCH_PAGES)]
    sem = asyncio.Semaphore(OCR_CONCURRENCY)

    async def run(batch: List[int]):
        async with sem:
            sub = await asyncio.to_thread(_pdf_subset, data, batch)
            result = await _ocr_call("application/pdf", sub, len(batch))
            pages, hw = _pages_from(result, len(batch))
            return dict(zip(batch, pages)), hw

    results = await asyncio.gather(*(run(b) for b in batches))
    merged: Dict[int, Tuple[str, str]] = {}
    handwritten = False
    for pages, hw in results:
        merged.update(pages)
        handwritten = handwritten or hw
    return merged, handwritten


def _illegible_share(text: str) -> float:
    words = re.findall(r"\[illegible(?: line)?\]|\S+", text)
    if not words:
        return 1.0
    bad = sum(1 for w in words if w.startswith("[illegible"))
    return bad / len(words)


_UNREADABLE_MSG = (
    "This document is too unclear to read reliably, so nothing was analysed "
    "rather than guessing at its contents. Try a sharper photo or scan: the "
    "whole page in frame, flat, in good light, without shadows."
)


def _judge(pages: List[Tuple[str, str]]) -> Tuple[bool, List[int]]:
    """(readable enough?, 1-based unclear pages) for OCR'd pages."""
    unclear = [i for i, (_, leg) in enumerate(pages, 1) if leg != "clear"]
    usable = [t for t, leg in pages if leg != "illegible" and _real_chars(t) >= 15]
    if not usable:
        return False, unclear
    if _illegible_share("\n".join(usable)) > MAX_ILLEGIBLE_SHARE:
        return False, unclear
    return True, unclear


def _ocr_note(ex: Extraction) -> str:
    """Prepended to OCR'd text so every analysis prompt sees it."""
    kind = "handwritten" if ex.handwritten else "scanned or photographed"
    if ex.method == "pdf-mixed":
        which = ", ".join(map(str, ex.ocr_pages))
        parts = [f"[Transcription note: page{'s' if len(ex.ocr_pages) > 1 else ''} "
                 f"{which} of this document {'are' if len(ex.ocr_pages) > 1 else 'is'} "
                 f"{kind} and {'were' if len(ex.ocr_pages) > 1 else 'was'} read by OCR."]
    else:
        parts = [f"[Transcription note: this document is {kind} and was read by OCR."]
    if ex.unclear_pages:
        where = ("page " + ", ".join(map(str, ex.unclear_pages))) if ex.pages > 1 else "the page"
        parts.append(f"Parts of {where} could not be read and are marked [illegible].")
    parts.append("Do not infer or fill in [illegible] passages; say they could not be read.]")
    return " ".join(parts)


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------

async def extract(data: bytes, content_type: str, filename: str = "") -> Extraction:
    kind = detect_kind(data, content_type, filename)
    if kind is None:
        raise ValueError(
            "That file type can't be read. Upload a PDF, Word document, "
            "photo (JPG, PNG, WEBP, HEIC) or text file."
        )

    if kind == "txt":
        return Extraction("ok", _clean(data.decode("utf-8", errors="replace")), "txt")

    if kind == "docx":
        try:
            return Extraction("ok", await asyncio.to_thread(_docx_text, data), "docx")
        except Exception as exc:
            raise ValueError(f"Could not read that Word file: {exc}")

    if kind == "doc":
        return Extraction("ok", await asyncio.to_thread(_doc_text, data), "doc")

    if kind.startswith("image/"):
        pages, hw = await _ocr_image(kind, data)
        ok, unclear = _judge(pages)
        if not ok:
            return Extraction("unreadable", method="image-ocr", ocr_pages=[1],
                              unclear_pages=unclear, handwritten=hw, message=_UNREADABLE_MSG)
        ex = Extraction("ok", _clean(pages[0][0]), "image-ocr", 1, [1], unclear, hw)
        ex.text = _ocr_note(ex) + "\n\n" + ex.text
        return ex

    # --- PDF: text layer where it exists, OCR where it doesn't -------------
    try:
        texts = await asyncio.to_thread(_pdf_pages, data)
    except Exception as exc:
        raise ValueError(f"Could not read that PDF: {exc}")
    if not texts:
        raise ValueError("That PDF has no pages.")

    need = [i for i, t in enumerate(texts, 1) if _real_chars(t) < MIN_TEXT_CHARS]
    if not need:
        return Extraction("ok", _join_pages(texts), "pdf-text", len(texts))

    skipped = need[MAX_OCR_PAGES:]
    need = need[:MAX_OCR_PAGES]
    ocr, hw = await _ocr_pdf_pages(data, need)

    ocr_list = [ocr[n] for n in need]
    ok, unclear_rel = _judge(ocr_list)
    unclear = [need[i - 1] for i in unclear_rel]
    all_scanned = len(need) + len(skipped) == len(texts)

    if not ok and all_scanned:
        return Extraction("unreadable", method="pdf-ocr", pages=len(texts), ocr_pages=need,
                          unclear_pages=unclear, handwritten=hw, message=_UNREADABLE_MSG)

    merged = list(texts)
    for n in need:
        text, leg = ocr[n]
        merged[n - 1] = text if leg != "illegible" else "[illegible page]"
    for n in skipped:
        merged[n - 1] = "[page not transcribed: over the OCR page limit]"

    ex = Extraction("ok", "", "pdf-ocr" if all_scanned else "pdf-mixed",
                    len(texts), need, unclear, hw)
    note = _ocr_note(ex)
    if skipped:
        note = note[:-1] + (f" Pages {skipped[0]}-{skipped[-1]} were beyond the "
                            f"{MAX_OCR_PAGES}-page OCR limit and were not read.]")
    ex.text = note + "\n\n" + _join_pages(merged)
    return ex


# ---------------------------------------------------------------------------
# Caching: memory -> in-flight prewarm -> saved sidecar -> extract
# ---------------------------------------------------------------------------

_CACHE_MAX = 48
_cache: "OrderedDict[str, Extraction]" = OrderedDict()
_inflight: Dict[str, "asyncio.Task[Extraction]"] = {}

_blob_client: Optional[httpx.AsyncClient] = None


def _client() -> httpx.AsyncClient:
    global _blob_client
    if _blob_client is None or _blob_client.is_closed:
        _blob_client = httpx.AsyncClient(
            timeout=60.0,
            limits=httpx.Limits(max_keepalive_connections=10, max_connections=20),
        )
    return _blob_client


async def close_client() -> None:
    global _blob_client
    if _blob_client is not None and not _blob_client.is_closed:
        await _blob_client.aclose()
    _blob_client = None


def _remember(path: str, ex: Extraction) -> None:
    _cache[path] = ex
    _cache.move_to_end(path)
    while len(_cache) > _CACHE_MAX:
        _cache.popitem(last=False)


def _sidecar(path: str) -> str:
    return f"{path}.extract.json"


async def _read_sidecar(path: str) -> Optional[Extraction]:
    try:
        url = await storage.signed_url(_sidecar(path), expires_in=60)
        resp = await _client().get(url)
        if resp.status_code != 200:
            return None
        raw = resp.json()
        if raw.get("version") != EXTRACT_VERSION:
            return None
        return Extraction(**{k: v for k, v in raw.items() if k in Extraction.__dataclass_fields__})
    except Exception:
        return None          # no sidecar yet (or unreachable): just extract


async def _write_sidecar(path: str, ex: Extraction) -> None:
    try:
        await storage.upload(
            _sidecar(path),
            json.dumps(asdict(ex), ensure_ascii=False).encode("utf-8"),
            "application/json",
        )
    except Exception as exc:
        logger.warning("Could not save extraction for %s (%s)", path, exc)


async def _extract_and_save(path: str, data: bytes, content_type: str, filename: str) -> Extraction:
    ex = await extract(data, content_type, filename)
    _remember(path, ex)
    await _write_sidecar(path, ex)
    if ex.ocr_pages:
        logger.info("OCR %s: method=%s pages=%s ocr=%s unclear=%s status=%s",
                    filename, ex.method, ex.pages, len(ex.ocr_pages), ex.unclear_pages, ex.status)
    return ex


def prewarm(path: str, data: bytes, content_type: str, filename: str) -> None:
    """Start extraction now, in the background, from the bytes just uploaded.

    Called by the upload endpoint after the file is stored. Never raises and
    never delays the upload response. A failure here is not an error: load()
    simply extracts again on first use and reports properly then.
    """
    if path in _cache or path in _inflight:
        return

    async def run() -> Extraction:
        try:
            return await _extract_and_save(path, data, content_type, filename)
        finally:
            _inflight.pop(path, None)

    task = asyncio.get_running_loop().create_task(run())
    task.add_done_callback(lambda t: t.cancelled() or t.exception())  # no "never retrieved" noise
    _inflight[path] = task


def _result(ex: Extraction) -> str:
    if ex.status == "unreadable":
        raise Unreadable(ex.message or _UNREADABLE_MSG)
    if not ex.text.strip():
        raise ValueError(
            "No readable text was found in this file. If it is a photo or "
            "scan, try a clearer one."
        )
    return ex.text


async def load(storage_path: str, content_type: str, filename: str) -> str:
    """Normalised text of a stored document. The single entry point.

    Raises ValueError (including Unreadable) with a user-facing message.
    """
    ex = _cache.get(storage_path)
    if ex is not None:
        _cache.move_to_end(storage_path)
        return _result(ex)

    task = _inflight.get(storage_path)
    if task is not None:
        try:
            return _result(await task)
        except (ValueError, Unreadable):
            raise
        except Exception as exc:
            logger.warning("Prewarm for %s failed (%s); retrying", storage_path, exc)

    ex = await _read_sidecar(storage_path)
    if ex is not None:
        _remember(storage_path, ex)
        return _result(ex)

    url = await storage.signed_url(storage_path, expires_in=120)
    blob = await _client().get(url)
    blob.raise_for_status()
    return _result(await _extract_and_save(storage_path, blob.content, content_type, filename))


async def forget(storage_path: str) -> None:
    """Drop the cached and saved extraction. Called when a document is deleted."""
    _cache.pop(storage_path, None)
    task = _inflight.pop(storage_path, None)
    if task is not None:
        task.cancel()
    await storage.delete(_sidecar(storage_path))